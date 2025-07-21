import os
import openpyxl
import subprocess
from datetime import datetime
import re
import logging

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

def validate_file_path(file_path, file_type):
    """Validate that a file exists and is a regular file"""
    if not os.path.exists(file_path):
        logger.error(f"{file_type} file not found: {file_path}")
        return False
    if not os.path.isfile(file_path):
        logger.error(f"{file_type} path is not a file: {file_path}")
        return False
    return True

def validate_directory_path(dir_path, dir_type):
    """Validate that a directory exists and is writable"""
    if not os.path.exists(dir_path):
        logger.error(f"{dir_type} directory not found: {dir_path}")
        return False
    if not os.path.isdir(dir_path):
        logger.error(f"{dir_type} path is not a directory: {dir_path}")
        return False
    if not os.access(dir_path, os.W_OK):
        logger.error(f"No write permission in {dir_type} directory: {dir_path}")
        return False
    return True

def get_timestamp_seconds(timestamp_str):
    """Convert timestamp string (HH:MM:SS or MM:SS) to seconds"""
    try:
        # Try to parse as HH:MM:SS
        dt = datetime.strptime(timestamp_str, "%H:%M:%S")
        return dt.hour * 3600 + dt.minute * 60 + dt.second
    except ValueError:
        # If that fails, try to parse as MM:SS
        try:
            dt = datetime.strptime(timestamp_str, "%M:%S")
            return dt.minute * 60 + dt.second
        except ValueError:
            logger.error(f"Invalid timestamp format: {timestamp_str}")
            return None

def extract_audio(input_file, output_file, start_time, end_time):
    """Extract audio segment using ffmpeg"""
    logger.info(f"\nExtracting audio from: {input_file}")
    logger.info(f"Output file: {output_file}")
    logger.info(f"Start time: {start_time} seconds")
    logger.info(f"End time: {end_time} seconds")
    
    # Validate input file
    if not validate_file_path(input_file, "Input"):
        return False
    
    # Validate output directory
    output_dir = os.path.dirname(output_file)
    if not validate_directory_path(output_dir, "Output"):
        return False
    
    # Build ffmpeg command string
    cmd = f"ffmpeg -i \"{input_file}\" -ss {start_time} -to {end_time} -c:a copy \"{output_file}\""
    try:
        logger.info(f"Running command: {cmd}")
        result = subprocess.run(cmd, shell=True, check=True, capture_output=True, text=True)
        logger.info(f"FFmpeg command completed successfully")
        logger.info(f"Output file exists: {os.path.exists(output_file)}")
        logger.info(f"Output file size: {os.path.getsize(output_file) if os.path.exists(output_file) else 'N/A'} bytes")
        return True
    except subprocess.CalledProcessError as e:
        logger.error(f"FFmpeg error: {e}")
        logger.error(f"Return code: {e.returncode}")
        logger.error(f"stderr: {e.stderr}")
        logger.error(f"stdout: {e.stdout}")
        logger.error(f"Full command: {cmd}")
        return False
    except Exception as e:
        logger.error(f"\nUnexpected error: {e}")
        return False

def extract_video(input_file, output_file, start_time, end_time):
    """Extract video segment using ffmpeg"""
    logger.info(f"\nExtracting video from: {input_file}")
    logger.info(f"Output file: {output_file}")
    logger.info(f"Start time: {start_time} seconds")
    logger.info(f"End time: {end_time} seconds")
    
    # Validate input file
    if not validate_file_path(input_file, "Input"):
        return False
    
    # Validate output directory
    output_dir = os.path.dirname(output_file)
    if not validate_directory_path(output_dir, "Output"):
        return False
    
    # Build ffmpeg command string with options to prevent slow motion
    cmd = f"ffmpeg -i \"{input_file}\" -ss {start_time - 3} -to {end_time + 3} -c:v copy -c:a copy -copyts -avoid_negative_ts make_zero \"{output_file}\""
    try:
        logger.info(f"Running command: {cmd}")
        result = subprocess.run(cmd, shell=True, check=True, capture_output=True, text=True)
        logger.info(f"FFmpeg command completed successfully")
        logger.info(f"Output file exists: {os.path.exists(output_file)}")
        logger.info(f"Output file size: {os.path.getsize(output_file) if os.path.exists(output_file) else 'N/A'} bytes")
        return True
    except subprocess.CalledProcessError as e:
        logger.error(f"FFmpeg error: {e}")
        logger.error(f"Return code: {e.returncode}")
        logger.error(f"stderr: {e.stderr}")
        logger.error(f"stdout: {e.stdout}")
        logger.error(f"Full command: {cmd}")
        return False
    except Exception as e:
        logger.error(f"\nUnexpected error: {e}")
        return False

def process_session_excel(excel_path, audio_dir, video_dir, audio_output_dir, video_output_dir):
    """Process a single session's Excel file"""
    logger.info(f"\nProcessing {excel_path}...")
    logger.info(f"Audio directory: {audio_dir}")
    logger.info(f"Video directory: {video_dir}")
    logger.info(f"Audio output directory: {audio_output_dir}")
    logger.info(f"Video output directory: {video_output_dir}")
    
    # Validate directories
    if not validate_directory_path(audio_dir, "Audio"):
        return False
    if not validate_directory_path(video_dir, "Video"):
        return False
    if not validate_directory_path(audio_output_dir, "Audio Output"):
        return False
    if not validate_directory_path(video_output_dir, "Video Output"):
        return False
    
    try:
        wb = openpyxl.load_workbook(excel_path)
        sheet = wb.active
        
        logger.info(f"\nSheet info for {excel_path}:")
        logger.info(f"Max row: {sheet.max_row}")
        logger.info(f"Max column: {sheet.max_column}")
        logger.info("Column headers:")
        for idx, cell in enumerate(sheet[1], 1):
            logger.info(f"Column {idx}: {cell.value}")
        
        current_dialog = None
        current_audio_file = None
        current_start_time = None
        current_end_time = None
        qd_count = 1
        cc_count = 1
        sr_count = 1
        
        for row in sheet.iter_rows(min_row=2):
            # Skip empty rows
            if not any(cell.value for cell in row):
                logger.debug(f"Skipping empty row")
                continue
                
            audio_file = str(row[2].value) if row[2].value is not None else None
            start_time = float(row[3].value) if row[3].value is not None else None
            end_time = float(row[4].value) if row[4].value is not None else None
            dialog_label = str(row[6].value) if row[6].value is not None else None
            
            # Skip rows with missing required fields
            if not audio_file:
                logger.warning(f"Skipping row - missing audio file")
                continue
            
            # Handle dialog labels
            if dialog_label and dialog_label.startswith("Start"):
                # Skip default dialog type
                if "default" in dialog_label.lower():
                    logger.info(f"Skipping default dialog label in row {row[0].row}")
                    continue
                
                # Reset previous dialog if any
                if current_dialog:
                    logger.warning(f"Found new Start label while processing {current_dialog}")
                    current_dialog = None
                    current_audio_file = None
                    current_start_time = None
                
                current_dialog = dialog_label.split()[1]
                current_audio_file = audio_file
                current_start_time = start_time
                logger.info(f"Found Start: {current_dialog} at {audio_file}")
                
                # Verify audio file exists
                audio_input = os.path.join(audio_dir, current_audio_file)
                logger.info(f"Audio input path: {audio_input}")
                if not validate_file_path(audio_input, "Audio"):
                    continue
                
                # Verify video file exists
                video_base_name = os.path.splitext(current_audio_file)[0]
                video_files = []
                for file in os.listdir(video_dir):
                    if file.startswith(video_base_name) and file.endswith('.MP4'):
                        video_files.append(file)
                
                if not video_files:
                    logger.warning(f"No video file found for {video_base_name}")
                    continue
                
                video_input = os.path.join(video_dir, video_files[0])
                logger.info(f"Found video file: {video_input}")
                logger.info(f"Video input path: {video_input}")
                if not validate_file_path(video_input, "Video"):
                    logger.warning(f"Video file not found: {video_input}")
                    continue
            
            elif dialog_label and dialog_label.startswith("End"):
                if not current_dialog:
                    logger.warning(f"Found End without matching Start: {dialog_label}")
                    continue
                
                if current_audio_file != audio_file:
                    logger.warning(f"Audio file mismatch: Start={current_audio_file}, End={audio_file}")
                    continue
                
                current_end_time = end_time
                
                # Process the dialog segment
                base_name = f"{current_audio_file}_"
                if "Question" in current_dialog:
                    base_name += f"qd_{qd_count}"
                    qd_count += 1
                elif "Chit-Chat" in current_dialog:
                    base_name += f"cc_{cc_count}"
                    cc_count += 1
                elif "Story" in current_dialog:
                    base_name += f"sr_{sr_count}"
                    sr_count += 1
                else:
                    logger.error(f"Unknown dialog type: {current_dialog}")
                    continue
                
                # Process audio
                audio_output = os.path.join(audio_output_dir, f"{base_name}.wav")
                logger.info(f"\nExtracting audio segment:")
                logger.info(f"Input: {audio_input}")
                logger.info(f"Output: {audio_output}")
                logger.info(f"Start: {current_start_time}")
                logger.info(f"End: {current_end_time}")
                result = extract_audio(audio_input, audio_output, current_start_time, current_end_time)
                if not result:
                    logger.error(f"Failed to extract audio for {base_name}")
                    continue
                
                # Process video
                video_output = os.path.join(video_output_dir, f"{base_name}.MP4")
                logger.info(f"\nExtracting video segment:")
                logger.info(f"Input: {video_input}")
                logger.info(f"Output: {video_output}")
                logger.info(f"Start: {current_start_time}")
                logger.info(f"End: {current_end_time}")
                result = extract_video(video_input, video_output, current_start_time, current_end_time)
                if not result:
                    logger.error(f"Failed to extract video for {base_name}")
                    continue
                
                logger.info(f"\nSuccessfully extracted {base_name}")
                logger.info(f"Audio output exists: {os.path.exists(audio_output)}")
                logger.info(f"Video output exists: {os.path.exists(video_output)}")
                logger.info(f"Audio output size: {os.path.getsize(audio_output) if os.path.exists(audio_output) else 'N/A'} bytes")
                logger.info(f"Video output size: {os.path.getsize(video_output) if os.path.exists(video_output) else 'N/A'} bytes")
                
                # Reset state for next dialog
                current_dialog = None
                current_audio_file = None
                current_start_time = None
                current_end_time = None
    except Exception as e:
        logger.error(f"Error processing Excel file: {str(e)}")
        return False
    return True

if __name__ == "__main__":
    # Input directories
    excel_dir = "/Users/mariyamohiuddin/Desktop/MIT Media Lab Projects/MoralQ Assets/session_1 xlsx/"
    audio_dir = "/Users/mariyamohiuddin/Desktop/MIT Media Lab Projects/MoralQ Assets/Session 1 Audio/"
    video_dir = "/Users/mariyamohiuddin/Desktop/MIT Media Lab Projects/MoralQ Assets/Session 1/"
    
    # Output directories
    audio_output_dir = "/Users/mariyamohiuddin/Desktop/MIT Media Lab Projects/MoralQ Assets/Child-Robot_Interaction_Clipping/Session 1 Audio/"
    video_output_dir = "/Users/mariyamohiuddin/Desktop/MIT Media Lab Projects/MoralQ Assets/Child-Robot_Interaction_Clipping/Session 1 Video/"
    
    # Create output directories with full permissions
    for dir_path in [audio_output_dir, video_output_dir]:
        try:
            os.makedirs(dir_path, exist_ok=True, mode=0o777)
        except Exception as e:
            logger.error(f"Error creating directory {dir_path}: {str(e)}")
            continue
    
    # Process all Excel files
    try:
        for excel_file in os.listdir(excel_dir):
            if not excel_file.endswith(".xlsx"):
                continue
                
            excel_path = os.path.join(excel_dir, excel_file)
            logger.info(f"\nProcessing {excel_file}...")
            logger.info(f"Audio output dir: {audio_output_dir}")
            logger.info(f"Video output dir: {video_output_dir}")
            logger.info(f"Audio output dir exists: {os.path.exists(audio_output_dir)}")
            logger.info(f"Video output dir exists: {os.path.exists(video_output_dir)}")
            logger.info(f"Audio output dir is dir: {os.path.isdir(audio_output_dir)}")
            logger.info(f"Video output dir is dir: {os.path.isdir(video_output_dir)}")
            logger.info(f"Audio directory: {audio_dir}")
            logger.info(f"Video directory: {video_dir}")
            
            result = process_session_excel(excel_path, audio_dir, video_dir, audio_output_dir, video_output_dir)
            if not result:
                logger.error(f"Failed to process {excel_file}")
                continue
    except Exception as e:
        logger.error(f"Error in main loop: {str(e)}")
